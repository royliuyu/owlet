"""Word extraction: headings, paragraphs, and tables become blocks."""

from datetime import datetime, timezone
from io import BytesIO

import pytest
from docx import Document

from core.connectors.base import RawDocument
from core.domain.errors import ParseError
from core.ingestion.parse import DocxParser, registered_formats
from core.ingestion.preview import parse_document, parse_to_chunks

VERSION = "test-docx"


def _docx() -> bytes:
    document = Document()
    document.add_heading("Installation", level=1)
    document.add_paragraph("Tighten the blade bolt before use.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Part"
    table.rows[0].cells[1].text = "126-8195"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_is_registered() -> None:
    assert "docx" in registered_formats()


def test_docx_parser_keeps_heading_paragraph_and_table() -> None:
    parsed = DocxParser().parse(_docx())

    assert [block.kind for block in parsed.blocks] == ["heading", "paragraph", "paragraph"]
    assert parsed.blocks[0].text == "Installation"
    assert "blade bolt" in parsed.blocks[1].text
    assert parsed.blocks[2].text == "Part | 126-8195"
    assert parsed.blocks[0].bboxes == []
    assert "Installation" in parsed.pages[0].text
    assert parsed.pages[0].page == 1
    assert [(block.kind, block.level, block.text, block.rows) for block in parsed.preview] == [
        ("heading", 1, "Installation", []),
        ("paragraph", 1, "Tighten the blade bolt before use.", []),
        ("table", 1, "", [["Part", "126-8195"]]),
    ]


def test_docx_preview_groups_lists_and_heading_levels() -> None:
    document = Document()
    document.add_heading("Installation", level=1)
    document.add_heading("Tools", level=2)
    document.add_paragraph("Check the guard", style="List Bullet")
    document.add_paragraph("Wear gloves", style="List Bullet")
    document.add_paragraph("First", style="List Number")
    buffer = BytesIO()
    document.save(buffer)

    preview = DocxParser().parse(buffer.getvalue()).preview

    assert [(block.kind, block.level, block.ordered, block.items) for block in preview] == [
        ("heading", 1, False, []),
        ("heading", 2, False, []),
        ("list", 1, False, ["Check the guard", "Wear gloves"]),
        ("list", 1, True, ["First"]),
    ]


def test_corrupt_docx_raises() -> None:
    with pytest.raises(ParseError):
        DocxParser().parse(b"not a word file")


def test_empty_docx_raises() -> None:
    document = Document()
    buffer = BytesIO()
    document.save(buffer)
    with pytest.raises(ParseError):
        DocxParser().parse(buffer.getvalue())


def test_parse_to_chunks_carries_the_heading() -> None:
    raw = RawDocument(
        source_id="lib/spec.docx",
        source="file",
        title="spec",
        uri="file:///spec.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=_docx(),
        content_hash="abc",
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
        extra={"format": "docx"},
    )

    chunks = parse_to_chunks(raw, chunking_version=VERSION)

    assert chunks
    assert chunks[0].doc_id == "file:lib/spec.docx"
    assert chunks[0].section_path == "Installation"
    assert "blade bolt" in chunks[0].text


def test_plain_text_decodes_as_one_page() -> None:
    raw = RawDocument(
        source_id="lib/todo.txt",
        source="note",
        title="todo",
        uri="file:///todo.txt",
        media_type="text/plain",
        content=b"buy milk\n\ncall Sam",
        content_hash="abc",
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
        extra={"format": "plain"},
    )

    parsed = parse_document(raw)

    assert parsed.pages[0].text == "buy milk\n\ncall Sam"
    assert parsed.blocks == []
