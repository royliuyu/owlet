"""Library, folder management, PDF preview, and chat event stream."""

from collections.abc import Callable
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from core.config import Settings
from core.domain.models import Chunk, Document
from core.interfaces.http.app import create_app
from core.interfaces.http.engine import Engine

MakePdf = Callable[[list[list[str]]], bytes]


@pytest.fixture
def papers(tmp_path: Path) -> Path:
    """The folder owlet is pointed at. The database lives outside it."""
    folder = tmp_path / "papers"
    folder.mkdir(exist_ok=True)
    return folder


def _app(tmp_path: Path, token: str = ""):
    return create_app(Settings(auth_token=token, data_dir=tmp_path / "state"))


def _client(tmp_path: Path, token: str = "") -> TestClient:
    folder = tmp_path / "papers"
    folder.mkdir(exist_ok=True)
    app = _app(tmp_path, token)
    app.state.engine.sources.add(folder, label="Papers", source_id="papers")
    return TestClient(app)


def test_lists_a_pdf_and_returns_extracted_pages(
    tmp_path: Path, papers: Path, make_pdf: MakePdf
) -> None:
    written = [["1 Introduction", "Attention is all you need."]]
    (papers / "attention.pdf").write_bytes(make_pdf(written))
    client = _client(tmp_path)

    collections = client.get("/api/v1/collections")
    assert collections.status_code == 200
    assert collections.json()[0]["id"] == "papers"
    assert collections.json()[0]["label"] == "Papers"
    assert collections.json()[0]["enabled"] is True

    files = client.get("/api/v1/collections/papers/files")
    assert files.status_code == 200
    body = files.json()
    assert len(body) == 1
    assert body[0]["id"] == "paper:papers/attention.pdf"
    assert body[0]["source"] == "paper"

    meta = client.get("/api/v1/documents/paper:papers/attention.pdf")
    assert meta.status_code == 200
    pages = meta.json()["pages"]
    assert pages[0]["page"] == 1
    assert "Attention is all you need." in pages[0]["text"]

    original = client.get("/api/v1/documents/paper:papers/attention.pdf?view=file")
    assert original.status_code == 200
    assert original.headers["content-type"].startswith("application/pdf")
    assert original.content.startswith(b"%PDF")


def test_docx_preview_returns_reading_blocks(tmp_path: Path, papers: Path) -> None:
    document = DocxDocument()
    document.add_heading("Installation", level=1)
    document.add_paragraph("Tighten the blade bolt before use.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Part"
    table.rows[0].cells[1].text = "126-8195"
    buffer = BytesIO()
    document.save(buffer)
    (papers / "spec.docx").write_bytes(buffer.getvalue())
    client = _client(tmp_path)

    meta = client.get("/api/v1/documents/file:papers/spec.docx")

    assert meta.status_code == 200
    blocks = meta.json()["blocks"]
    assert [block["kind"] for block in blocks] == ["heading", "paragraph", "table"]
    assert blocks[0]["text"] == "Installation"
    assert blocks[0]["level"] == 1
    assert blocks[2]["rows"] == [["Part", "126-8195"]]


def test_a_folder_is_added_and_removed_without_a_restart(tmp_path: Path, papers: Path) -> None:
    client = TestClient(_app(tmp_path))
    assert client.get("/api/v1/collections").json() == []

    created = client.post("/api/v1/collections", json={"path": str(papers), "label": "Papers"})
    assert created.status_code == 201
    assert created.json() == {
        "id": "papers",
        "label": "Papers",
        "path": str(papers.resolve()),
        "enabled": True,
    }
    assert [item["id"] for item in client.get("/api/v1/collections").json()] == ["papers"]

    assert client.delete("/api/v1/collections/papers").status_code == 204
    assert client.get("/api/v1/collections").json() == []
    assert client.delete("/api/v1/collections/papers").status_code == 404


def test_removing_a_folder_drops_its_index_and_leaves_the_files(
    tmp_path: Path, papers: Path
) -> None:
    notes = tmp_path / "notes"
    notes.mkdir()
    (papers / "attention.pdf").write_bytes(b"%PDF-1.4 leftover")
    client = TestClient(_app(tmp_path))
    assert client.post(
        "/api/v1/collections", json={"path": str(papers), "label": "Papers", "id": "papers"}
    ).status_code == 201
    assert client.post(
        "/api/v1/collections", json={"path": str(notes), "label": "Notes", "id": "notes"}
    ).status_code == 201
    engine: Engine = client.app.state.engine
    _passage(engine, "paper:papers/attention.pdf", "papers", "attention is all you need")
    _passage(engine, "note:notes/lab.md", "notes", "the lab notebook stays")
    paper_vector = [1.0, *([0.0] * (engine.vectors.dim - 1))]
    note_vector = [0.0, *([0.0] * (engine.vectors.dim - 2)), 1.0]
    engine.vectors.upsert(
        ["paper:papers/attention.pdf:0", "note:notes/lab.md:0"],
        [paper_vector, note_vector],
    )
    assert engine.documents.stats() == {"documents": 2, "chunks": 2, "embedded": 2}

    removed = client.delete("/api/v1/collections/papers")

    assert removed.status_code == 204
    assert [item["id"] for item in client.get("/api/v1/collections").json()] == ["notes"]
    assert (papers / "attention.pdf").is_file()
    assert engine.documents.stats() == {"documents": 1, "chunks": 1, "embedded": 1}
    assert list(engine.documents.indexed()) == ["note:notes/lab.md"]
    assert engine.documents.search_fulltext("attention", limit=5) == []
    assert [hit.chunk_id for hit in engine.vectors.search(paper_vector, limit=5)] == [
        "note:notes/lab.md:0"
    ]
    assert [hit.chunk_id for hit in engine.vectors.search(note_vector, limit=5)] == [
        "note:notes/lab.md:0"
    ]
    status = client.get("/api/v1/index/status").json()
    assert status["running"] is False
    assert status["documents"] == 1


def test_a_folder_is_not_removed_while_indexing(tmp_path: Path, papers: Path) -> None:
    client = _client(tmp_path)
    engine: Engine = client.app.state.engine
    _passage(engine, "paper:papers/attention.pdf", "papers", "attention is all you need")

    class _Running:
        def done(self) -> bool:
            return False

    engine._task = _Running()  # type: ignore[assignment]
    refused = client.delete("/api/v1/collections/papers")
    assert refused.status_code == 409
    assert "index" in refused.json()["detail"].lower()
    assert engine.documents.stats()["documents"] == 1
    assert [item["id"] for item in client.get("/api/v1/collections").json()] == ["papers"]


def test_browse_lists_child_folders_and_starts_at_the_drives(tmp_path: Path) -> None:
    (tmp_path / "papers").mkdir()
    (tmp_path / "notes").mkdir()
    (tmp_path / "readme.txt").write_text("x", encoding="utf-8")
    (tmp_path / "papers" / "2024").mkdir()
    client = TestClient(_app(tmp_path))

    listed = client.get("/api/v1/browse", params={"path": str(tmp_path)})
    assert listed.status_code == 200
    body = listed.json()
    assert body["path"] == str(tmp_path.resolve())
    assert {item["name"] for item in body["entries"]} == {"notes", "papers"}

    into = client.get("/api/v1/browse", params={"path": str(tmp_path / "papers")})
    assert into.status_code == 200
    assert into.json()["parent"] == str((tmp_path / "papers").resolve().parent)
    assert [item["name"] for item in into.json()["entries"]] == ["2024"]

    missing = client.get("/api/v1/browse", params={"path": str(tmp_path / "nowhere")})
    assert missing.status_code == 404
    note = tmp_path / "readme.txt"
    assert client.get("/api/v1/browse", params={"path": str(note)}).status_code == 404

    roots = client.get("/api/v1/browse")
    assert roots.status_code == 200
    drives = roots.json()
    assert drives["parent"] is None
    assert drives["entries"]
    assert all(Path(item["path"]).is_dir() for item in drives["entries"])

    top = client.get("/api/v1/browse", params={"path": tmp_path.anchor})
    assert top.status_code == 200
    assert top.json()["parent"] == ("" if tmp_path.anchor != "/" else None)


def test_a_folder_that_is_not_there_is_refused(tmp_path: Path, papers: Path) -> None:
    client = TestClient(_app(tmp_path))
    missing = client.post("/api/v1/collections", json={"path": str(tmp_path / "nowhere")})
    assert missing.status_code == 400
    assert "folder" in missing.json()["detail"]

    assert client.post("/api/v1/collections", json={"path": str(papers)}).status_code == 201
    again = client.post("/api/v1/collections", json={"path": str(papers)})
    assert again.status_code == 409


def test_a_disabled_folder_keeps_its_row_but_serves_no_files(
    tmp_path: Path, papers: Path, make_pdf: MakePdf
) -> None:
    (papers / "attention.pdf").write_bytes(make_pdf([["Hello"]]))
    client = _client(tmp_path)
    assert len(client.get("/api/v1/collections/papers/files").json()) == 1

    patched = client.patch("/api/v1/collections/papers", json={"enabled": False})
    assert patched.status_code == 200
    assert patched.json()["enabled"] is False
    assert client.get("/api/v1/collections/papers/files").json() == []
    assert client.get("/api/v1/documents/paper:papers/attention.pdf").status_code == 404


def test_unknown_collection_and_bad_document_id(tmp_path: Path) -> None:
    client = _client(tmp_path)
    assert client.get("/api/v1/collections/missing/files").status_code == 404
    assert client.get("/api/v1/documents/not-an-id").status_code == 400
    assert client.get("/api/v1/documents/paper:other/secret.pdf").status_code == 400


def test_broken_pdf_is_rejected(tmp_path: Path, papers: Path) -> None:
    (papers / "broken.pdf").write_bytes(b"not a pdf")
    client = _client(tmp_path)
    response = client.get("/api/v1/documents/paper:papers/broken.pdf")
    assert response.status_code == 422
    assert "PDF" in response.json()["message"]


def test_chat_lists_every_document_filed_under_a_person(tmp_path: Path) -> None:
    client = _client(tmp_path)
    engine: Engine = client.app.state.engine
    _record(
        engine,
        "paper:arod",
        "AROD: Adaptive Real-Time Object Detection",
        ("Yu Liu", "Kyoung-Don Kang"),
        2024,
    )
    _record(engine, "paper:prep", "Preprocessing via Deep Learning", ("Yu Liu",), 2023)
    _record(engine, "paper:frames", "Filtering empty video frames", ("Yu Liu",), 2024)
    _record(
        engine,
        "paper:corun",
        "Corun: Concurrent Inference",
        ("Yu Liu", "Anurag Andhare"),
        2024,
    )
    _record(engine, "paper:hadd", "HADD: High-Accuracy Drift Detection", ("Yu Liu",), 2022)
    _record(engine, "paper:survey", "Survey of edge devices", ("Ada Lovelace",), 2021)

    listed = client.post(
        "/api/v1/chat",
        json={
            "question": "you search the library, and list all the paper title written by Yu Liu",
            "focus_doc_id": "paper:arod",
        },
    )

    assert listed.status_code == 200
    assert "list_records" in listed.text
    assert "event: error" not in listed.text
    assert listed.text.count("event: citation") == 5
    for title in (
        "AROD: Adaptive Real-Time Object Detection",
        "Preprocessing via Deep Learning",
        "Filtering empty video frames",
        "Corun: Concurrent Inference",
        "HADD: High-Accuracy Drift Detection",
    ):
        assert title in listed.text
    assert "Survey of edge devices" not in listed.text
    assert "Yu Liu is listed on 5 documents" in listed.text

    one = client.post("/api/v1/chat", json={"question": "Who wrote AROD?"})
    assert "list_records" not in one.text
    assert "AROD: Adaptive Real-Time Object Detection is by Yu Liu, Kyoung-Don Kang." in one.text


def _passage(engine: Engine, doc_id: str, root_id: str, text: str) -> None:
    now = datetime.now(tz=timezone.utc)
    source, source_id = doc_id.split(":", 1)
    engine.documents.upsert_document(
        Document(
            id=doc_id,
            source=source,  # type: ignore[arg-type]
            source_id=source_id,
            title=source_id,
            uri=f"file:///{source_id}",
            created_at=now,
            updated_at=now,
            content_hash=doc_id,
        ),
        root_id=root_id,
        page_count=1,
        parser="pdf",
    )
    engine.documents.replace_chunks(
        doc_id,
        [
            Chunk(
                id=f"{doc_id}:0",
                doc_id=doc_id,
                ord=0,
                text=text,
                token_count=4,
                page_start=1,
                page_end=1,
            )
        ],
        chunking_version="test",
    )


def _record(
    engine: Engine,
    doc_id: str,
    title: str,
    authors: tuple[str, ...],
    year: int,
) -> None:
    now = datetime.now(tz=timezone.utc)
    source, source_id = doc_id.split(":", 1)
    engine.documents.upsert_document(
        Document(
            id=doc_id,
            source=source,  # type: ignore[arg-type]
            source_id=source_id,
            title=title,
            uri=f"file:///{source_id}",
            created_at=now,
            updated_at=now,
            content_hash=doc_id,
        ),
        root_id="papers",
        page_count=1,
        parser="pdf",
    )
    engine.documents.set_bibliography(
        doc_id,
        {"v": 1, "title": title, "authors": list(authors), "year": year, "venue": "", "doi": ""},
    )


def test_chat_says_what_is_missing_when_nothing_is_indexed(tmp_path: Path) -> None:
    client = _client(tmp_path)
    response = client.post("/api/v1/chat", json={"question": "What is attention?"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/event-stream")
    assert "event: start" in response.text
    assert "event: tool" in response.text
    assert "event: error" in response.text
    assert "Nothing is indexed yet" in response.text
    assert "event: done" in response.text
    assert client.post("/api/v1/chat", json={"question": ""}).status_code == 422


def test_index_status_starts_empty(tmp_path: Path) -> None:
    client = _client(tmp_path)
    status = client.get("/api/v1/index/status").json()
    assert status["running"] is False
    assert status["documents"] == 0
    assert status["chunks"] == 0
    assert client.post("/api/v1/search", json={"query": "attention"}).json() == []
    assert client.post("/api/v1/index/cancel").json()["running"] is False


def test_token_is_required_for_api_calls(tmp_path: Path) -> None:
    client = _client(tmp_path, token="secret")
    assert client.get("/api/v1/health").status_code == 200
    assert client.get("/api/v1/collections").status_code == 401
    assert client.post("/api/v1/collections", json={"path": str(tmp_path)}).status_code == 401
    assert client.post("/api/v1/session", json={"token": "nope"}).status_code == 401

    signed_in = client.post("/api/v1/session", json={"token": "secret"})
    assert signed_in.status_code == 200
    assert client.get("/api/v1/collections").status_code == 200


def test_confirm_action_does_not_run(tmp_path: Path) -> None:
    client = _client(tmp_path)
    response = client.post("/api/v1/actions/a1/confirm")
    assert response.status_code == 501
    assert "a1" in response.json()["message"]
